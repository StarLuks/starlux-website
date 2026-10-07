"""Справочники: типы цен, группы номенклатуры, карточки товаров, изображения, импорт из Excel."""
import base64
import io
import os
import secrets
from decimal import Decimal, InvalidOperation

import boto3

PRODUCT_FIELDS = {
    'name': 'name', 'fullName': 'full_name', 'article': 'article', 'code1c': 'code_1c',
    'barcode': 'barcode', 'unit': 'unit', 'manufacturer': 'manufacturer', 'dimensions': 'dimensions',
    'pack': 'pack',
}


def _q(cur, sql, args=None):
    cur.execute(sql.replace('{S}', os.environ.get('MAIN_DB_SCHEMA', 't_p16770056_starlux_website')), args)


def _s3():
    return boto3.client('s3', endpoint_url='https://bucket.poehali.dev',
                        aws_access_key_id=os.environ['AWS_ACCESS_KEY_ID'],
                        aws_secret_access_key=os.environ['AWS_SECRET_ACCESS_KEY'])


def _num(v, default=None):
    if v is None or (isinstance(v, str) and not v.strip()):
        return default
    try:
        return Decimal(str(v).replace(',', '.').replace(' ', '').replace('\xa0', ''))
    except InvalidOperation:
        return default


def _sync_main_price(cur, product_id=None):
    where = "AND p.id = %s" if product_id else ""
    _q(cur, "UPDATE {S}.products p SET price = pp.price, updated_at = NOW() FROM {S}.product_prices pp "
            "JOIN {S}.price_types t ON t.id = pp.price_type_id AND t.is_main "
            "WHERE pp.product_id = p.id " + where, (product_id,) if product_id else None)


def list_price_types(cur):
    _q(cur, "SELECT t.id, t.name, t.code_1c AS \"code1c\", t.is_main AS \"isMain\", t.active, "
            "(SELECT COUNT(*) FROM {S}.product_prices pp WHERE pp.price_type_id = t.id) AS \"pricesCount\" "
            "FROM {S}.price_types t ORDER BY t.sort, t.id")
    return cur.fetchall()


def list_groups(cur):
    _q(cur, "SELECT g.id, g.name, g.code_1c AS \"code1c\", "
            "(SELECT COUNT(*) FROM {S}.products p WHERE p.group_id = g.id) AS \"productsCount\" "
            "FROM {S}.product_groups g ORDER BY g.sort, g.name")
    return cur.fetchall()


def list_products(cur):
    _q(cur, "SELECT p.id, p.group_id AS \"groupId\", p.active, p.name, p.full_name AS \"fullName\", p.article, "
            "p.code_1c AS \"code1c\", p.barcode, p.unit, p.manufacturer, p.dimensions, p.pack_kg AS \"weight\", "
            "p.pack, p.stock, p.updated_at AS \"updatedAt\" FROM {S}.products p ORDER BY p.sort, p.name")
    rows = cur.fetchall()
    _q(cur, "SELECT id, product_id, url, is_main AS \"isMain\" FROM {S}.product_images ORDER BY is_main DESC, sort, id")
    imgs = {}
    for i in cur.fetchall():
        imgs.setdefault(i.pop('product_id'), []).append(i)
    _q(cur, "SELECT product_id, price_type_id, price FROM {S}.product_prices")
    prices = {}
    for pr in cur.fetchall():
        prices.setdefault(pr['product_id'], {})[str(pr['price_type_id'])] = pr['price']
    for r in rows:
        r['images'] = imgs.get(r['id'], [])
        r['prices'] = prices.get(r['id'], {})
    return rows


def _save_prices(cur, pid, prices):
    for tid, val in (prices or {}).items():
        v = _num(val)
        if v is None:
            _q(cur, "DELETE FROM {S}.product_prices WHERE product_id = %s AND price_type_id = %s", (pid, int(tid)))
        else:
            _q(cur, "INSERT INTO {S}.product_prices (product_id, price_type_id, price) VALUES (%s,%s,%s) "
                    "ON CONFLICT (product_id, price_type_id) DO UPDATE SET price = EXCLUDED.price, updated_at = NOW()",
               (pid, int(tid), v))
    _sync_main_price(cur, pid)


def _upsert_product(cur, pid, data, group_id):
    vals = {col: str(data.get(key) or '').strip() for key, col in PRODUCT_FIELDS.items() if key in data}
    if 'weight' in data:
        vals['pack_kg'] = _num(data.get('weight'), Decimal(1))
    if 'active' in data:
        vals['active'] = bool(data['active'])
    if 'stock' in data and data.get('stock') not in (None, ''):
        vals['stock'] = int(_num(data.get('stock'), Decimal(0)))
    if group_id is not None:
        vals['group_id'] = group_id
        _q(cur, "SELECT name FROM {S}.product_groups WHERE id = %s", (group_id,))
        g = cur.fetchone()
        vals['category'] = g['name'] if g else 'Прочее'
    if pid:
        if vals:
            cols = ', '.join(f"{c} = %s" for c in vals)
            _q(cur, "UPDATE {S}.products SET " + cols + ", updated_at = NOW() WHERE id = %s", (*vals.values(), pid))
        return pid
    pid = 'n' + secrets.token_hex(5)
    vals.setdefault('category', 'Прочее')
    vals.setdefault('name', 'Без названия')
    _q(cur, "SELECT COALESCE(MAX(sort), 0) + 1 AS s FROM {S}.products")
    vals['sort'] = cur.fetchone()['s']
    cols = ', '.join(['id', *vals.keys()])
    ph = ', '.join(['%s'] * (len(vals) + 1))
    _q(cur, "INSERT INTO {S}.products (" + cols + ") VALUES (" + ph + ")", (pid, *vals.values()))
    return pid


def handle(cur, action, body):
    """Возвращает (код, данные) или None, если действие не относится к справочникам."""
    if action == 'price_types':
        return 200, {'priceTypes': list_price_types(cur)}

    if action == 'price_type_save':
        name = (body.get('name') or '').strip()
        if not name:
            return 400, {'error': 'Укажите наименование типа цены'}
        tid = body.get('id')
        code = (body.get('code1c') or '').strip()
        active = bool(body.get('active', True))
        if tid:
            _q(cur, "UPDATE {S}.price_types SET name = %s, code_1c = %s, active = %s WHERE id = %s",
               (name, code, active, int(tid)))
        else:
            _q(cur, "INSERT INTO {S}.price_types (name, code_1c, active, sort) "
                    "VALUES (%s, %s, %s, (SELECT COALESCE(MAX(sort), 0) + 1 FROM {S}.price_types)) RETURNING id",
               (name, code, active))
            tid = cur.fetchone()['id']
        if body.get('isMain'):
            _q(cur, "UPDATE {S}.price_types SET is_main = (id = %s), active = active OR id = %s", (int(tid), int(tid)))
            _sync_main_price(cur)
        _q(cur, "SELECT is_main, active FROM {S}.price_types WHERE id = %s", (int(tid),))
        r = cur.fetchone()
        if r['is_main'] and not r['active']:
            _q(cur, "UPDATE {S}.price_types SET active = TRUE WHERE id = %s", (int(tid),))
            return 200, {'ok': True, 'id': tid, 'warning': 'Основной тип цены нельзя отключить'}
        return 200, {'ok': True, 'id': tid}

    if action == 'price_type_delete':
        tid = int(body.get('id', 0))
        _q(cur, "SELECT is_main FROM {S}.price_types WHERE id = %s", (tid,))
        r = cur.fetchone()
        if not r:
            return 404, {'error': 'Тип цены не найден'}
        if r['is_main']:
            return 409, {'error': 'Нельзя удалить основной тип цены — сначала назначьте основным другой'}
        _q(cur, "DELETE FROM {S}.product_prices WHERE price_type_id = %s", (tid,))
        _q(cur, "DELETE FROM {S}.price_types WHERE id = %s", (tid,))
        return 200, {'ok': True}

    if action == 'groups':
        return 200, {'groups': list_groups(cur)}

    if action == 'group_save':
        name = (body.get('name') or '').strip()
        if not name:
            return 400, {'error': 'Укажите наименование группы'}
        code = (body.get('code1c') or '').strip()
        if body.get('id'):
            gid = int(body['id'])
            _q(cur, "UPDATE {S}.product_groups SET name = %s, code_1c = %s WHERE id = %s", (name, code, gid))
            _q(cur, "UPDATE {S}.products SET category = %s WHERE group_id = %s", (name, gid))
        else:
            _q(cur, "INSERT INTO {S}.product_groups (name, code_1c, sort) "
                    "VALUES (%s, %s, (SELECT COALESCE(MAX(sort), 0) + 1 FROM {S}.product_groups)) RETURNING id", (name, code))
            gid = cur.fetchone()['id']
        return 200, {'ok': True, 'id': gid}

    if action == 'group_delete':
        gid = int(body.get('id', 0))
        _q(cur, "SELECT COUNT(*) AS c FROM {S}.products WHERE group_id = %s", (gid,))
        if cur.fetchone()['c']:
            return 409, {'error': 'В группе есть товары — перенесите их в другую группу перед удалением'}
        _q(cur, "DELETE FROM {S}.product_groups WHERE id = %s", (gid,))
        return 200, {'ok': True}

    if action == 'nomenclature':
        return 200, {'products': list_products(cur), 'groups': list_groups(cur), 'priceTypes': list_price_types(cur)}

    if action == 'product_save':
        if not (body.get('name') or '').strip():
            return 400, {'error': 'Укажите наименование товара'}
        gid = body.get('groupId')
        if not gid:
            return 400, {'error': 'Выберите группу номенклатуры'}
        pid = _upsert_product(cur, body.get('id'), body, int(gid))
        _save_prices(cur, pid, body.get('prices'))
        return 200, {'ok': True, 'id': pid}

    if action == 'product_toggle':
        _q(cur, "UPDATE {S}.products SET active = NOT active, updated_at = NOW() WHERE id = %s RETURNING active",
           (str(body.get('id')),))
        r = cur.fetchone()
        if not r:
            return 404, {'error': 'Товар не найден'}
        return 200, {'active': r['active']}

    if action == 'image_upload':
        pid = str(body.get('productId') or '')
        _q(cur, "SELECT id FROM {S}.products WHERE id = %s", (pid,))
        if not cur.fetchone():
            return 404, {'error': 'Товар не найден'}
        raw = body.get('data') or ''
        if ',' in raw[:100]:
            raw = raw.split(',', 1)[1]
        data = base64.b64decode(raw)
        ctype = body.get('contentType') or 'image/jpeg'
        ext = {'image/png': 'png', 'image/webp': 'webp'}.get(ctype, 'jpg')
        key = f"products/{pid}/{secrets.token_hex(8)}.{ext}"
        _s3().put_object(Bucket='files', Key=key, Body=data, ContentType=ctype)
        url = f"https://cdn.poehali.dev/projects/{os.environ['AWS_ACCESS_KEY_ID']}/bucket/{key}"
        _q(cur, "SELECT COUNT(*) AS c FROM {S}.product_images WHERE product_id = %s", (pid,))
        first = cur.fetchone()['c'] == 0
        _q(cur, "INSERT INTO {S}.product_images (product_id, url, is_main, sort) VALUES (%s, %s, %s, "
                "(SELECT COALESCE(MAX(sort), 0) + 1 FROM {S}.product_images WHERE product_id = %s)) "
                "RETURNING id, url, is_main AS \"isMain\"", (pid, url, first, pid))
        return 200, {'image': cur.fetchone()}

    if action == 'image_delete':
        _q(cur, "DELETE FROM {S}.product_images WHERE id = %s RETURNING product_id, is_main", (int(body.get('id', 0)),))
        r = cur.fetchone()
        if r and r['is_main']:
            _q(cur, "UPDATE {S}.product_images SET is_main = TRUE WHERE id = (SELECT id FROM {S}.product_images "
                    "WHERE product_id = %s ORDER BY sort, id LIMIT 1)", (r['product_id'],))
        return 200, {'ok': True}

    if action == 'image_main':
        _q(cur, "SELECT product_id FROM {S}.product_images WHERE id = %s", (int(body.get('id', 0)),))
        r = cur.fetchone()
        if not r:
            return 404, {'error': 'Изображение не найдено'}
        _q(cur, "UPDATE {S}.product_images SET is_main = (id = %s) WHERE product_id = %s",
           (int(body['id']), r['product_id']))
        return 200, {'ok': True}

    if action == 'import_template':
        return 200, {'file': _template(cur), 'name': 'Шаблон_номенклатуры.xlsx'}

    if action == 'import_excel':
        return _import_excel(cur, body.get('file') or '')

    return None


BASE_COLS = ['Код 1С', 'Артикул', 'Наименование', 'Полное наименование', 'Группа', 'Код группы 1С',
             'Единица измерения', 'Вид упаковки', 'Масса, кг', 'Габариты', 'Штрихкод', 'Производитель',
             'Остаток', 'Активна']


def _template(cur):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    types = [t for t in list_price_types(cur) if t['active']]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Номенклатура'
    header = BASE_COLS + [f"Цена: {t['name']}" for t in types]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1D4F9C')
    for p in list_products(cur):
        _q(cur, "SELECT name, code_1c FROM {S}.product_groups WHERE id = %s", (p['groupId'],))
        g = cur.fetchone() or {'name': '', 'code_1c': ''}
        ws.append([p['code1c'], p['article'], p['name'], p['fullName'], g['name'], g['code_1c'], p['unit'], p['pack'],
                   float(p['weight']), p['dimensions'], p['barcode'], p['manufacturer'], p['stock'],
                   'да' if p['active'] else 'нет']
                  + [float(p['prices'][str(t['id'])]) if str(t['id']) in p['prices'] else None for t in types])
    for i, h in enumerate(header):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(12, min(40, len(h) + 4))
    ws.column_dimensions['C'].width = 36
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


TEXT_COLS = {'name': 'name', 'full_name': 'fullName', 'article': 'article', 'code_1c': 'code1c',
             'barcode': 'barcode', 'unit': 'unit', 'manufacturer': 'manufacturer', 'dimensions': 'dimensions',
             'pack': 'pack'}
UPD_COLS = ['name', 'full_name', 'article', 'code_1c', 'barcode', 'unit', 'manufacturer', 'dimensions', 'pack',
            'pack_kg', 'stock', 'active', 'group_id', 'category']
UPD_TYPES = {**{c: 'text' for c in UPD_COLS}, 'pack_kg': 'numeric', 'stock': 'integer', 'active': 'boolean',
             'group_id': 'integer'}


def _import_excel(cur, file_b64):
    from openpyxl import load_workbook
    if ',' in file_b64[:100]:
        file_b64 = file_b64.split(',', 1)[1]
    try:
        wb = load_workbook(io.BytesIO(base64.b64decode(file_b64)), read_only=True, data_only=True)
    except Exception:
        return 400, {'error': 'Не удалось прочитать файл. Загрузите файл Excel в формате .xlsx'}
    rows = list(wb.active.iter_rows(values_only=True))
    if len(rows) < 2:
        return 400, {'error': 'В файле нет строк с товарами'}
    head = [str(h or '').strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            if n.lower() in head:
                return head.index(n.lower())
        return None

    idx = {
        'code1c': col('Код 1С', 'Код1С', 'Код'), 'article': col('Артикул'), 'name': col('Наименование'),
        'fullName': col('Полное наименование'), 'group': col('Группа', 'Группа номенклатуры'),
        'groupCode': col('Код группы 1С', 'Код группы'), 'unit': col('Единица измерения', 'Ед. изм.', 'Ед.изм.'),
        'pack': col('Вид упаковки', 'Упаковка'), 'weight': col('Масса, кг', 'Масса', 'Вес'),
        'dimensions': col('Габариты'), 'barcode': col('Штрихкод'), 'manufacturer': col('Производитель'),
        'stock': col('Остаток'), 'active': col('Активна', 'Активен'),
    }
    if idx['name'] is None:
        return 400, {'error': 'В первой строке не найдена колонка «Наименование». Скачайте шаблон и заполните его'}

    _q(cur, "SELECT id, lower(name) AS name FROM {S}.price_types")
    ptypes = {r['name']: r['id'] for r in cur.fetchall()}
    price_cols = {}
    for i, h in enumerate(head):
        if h.startswith('цена:') and h[5:].strip() in ptypes:
            price_cols[i] = ptypes[h[5:].strip()]

    _q(cur, "SELECT id, name, code_1c FROM {S}.product_groups")
    groups = cur.fetchall()
    gname_by_id = {g['id']: g['name'] for g in groups}
    by_gcode = {g['code_1c']: g['id'] for g in groups if g['code_1c']}
    by_gname = {g['name'].lower(): g['id'] for g in groups}
    _q(cur, "SELECT id, code_1c, article, lower(name) AS name FROM {S}.products")
    prods = cur.fetchall()
    by_code = {p['code_1c']: p['id'] for p in prods if p['code_1c']}
    by_art = {p['article']: p['id'] for p in prods if p['article']}
    by_name = {p['name']: p['id'] for p in prods}
    existing = {p['id'] for p in prods}

    new_rows, upd_rows, price_map = {}, {}, {}
    errors = []
    for n, r in enumerate(rows[1:], start=2):
        def v(k):
            i = idx[k]
            if i is None or i >= len(r) or r[i] is None:
                return None
            x = r[i]
            if isinstance(x, float) and x.is_integer() and k in ('code1c', 'article', 'barcode', 'groupCode'):
                x = int(x)
            x = str(x).strip()
            return x if x != '' else None

        name = v('name')
        if not name:
            if any(c not in (None, '') for c in r):
                errors.append(f"Строка {n}: не указано наименование")
            continue
        gname, gcode = v('group'), v('groupCode')
        gid = (by_gcode.get(gcode) if gcode else None) or (by_gname.get(gname.lower()) if gname else None)
        if not gid and (gname or gcode):
            title = gname or f"Группа {gcode}"
            _q(cur, "INSERT INTO {S}.product_groups (name, code_1c, sort) VALUES (%s, %s, "
                    "(SELECT COALESCE(MAX(sort), 0) + 1 FROM {S}.product_groups)) RETURNING id",
               (title, gcode or ''))
            gid = cur.fetchone()['id']
            gname_by_id[gid] = title
            if gcode:
                by_gcode[gcode] = gid
            by_gname[title.lower()] = gid
        code, art = v('code1c'), v('article')
        pid = (by_code.get(code) if code else None) or (by_art.get(art) if art else None) or by_name.get(name.lower())
        if not pid and not gid:
            errors.append(f"Строка {n} ({name}): не указана группа для нового товара")
            continue

        rec = {c: v(k) for c, k in TEXT_COLS.items()}
        weight, stock = v('weight'), v('stock')
        rec['pack_kg'] = _num(weight) if weight is not None else None
        if weight is not None and rec['pack_kg'] is None:
            errors.append(f"Строка {n} ({name}): некорректная масса «{weight}»")
        st = _num(stock) if stock is not None else None
        if stock is not None and st is None:
            errors.append(f"Строка {n} ({name}): некорректный остаток «{stock}»")
        rec['stock'] = int(st) if st is not None else None
        act = v('active')
        rec['active'] = None if act is None else act.lower() not in ('нет', 'no', '0', 'false', 'ложь', '-')
        rec['group_id'] = gid
        rec['category'] = gname_by_id.get(gid) if gid else None

        if not pid:
            pid = 'n' + secrets.token_hex(5)
        target = upd_rows if pid in existing else new_rows
        prev = target.get(pid)
        target[pid] = {**prev, **{k: x for k, x in rec.items() if x is not None}} if prev else rec
        if code:
            by_code[code] = pid
        if art:
            by_art[art] = pid
        by_name[name.lower()] = pid

        for i, tid in price_cols.items():
            if i < len(r) and r[i] not in (None, ''):
                val = _num(r[i])
                if val is None:
                    errors.append(f"Строка {n} ({name}): некорректная цена «{r[i]}»")
                else:
                    price_map[(pid, tid)] = round(val, 2)

    from psycopg2.extras import execute_values
    schema = os.environ.get('MAIN_DB_SCHEMA', 't_p16770056_starlux_website')

    if new_rows:
        _q(cur, "SELECT COALESCE(MAX(sort), 0) AS s FROM {S}.products")
        sort = cur.fetchone()['s']
        data = []
        for pid, rec in new_rows.items():
            sort += 1
            data.append((
                pid, rec['name'], rec.get('full_name') or rec['name'], rec.get('article') or '',
                rec.get('code_1c') or '', rec.get('barcode') or '', rec.get('unit') or 'кор.',
                rec.get('manufacturer') or '', rec.get('dimensions') or '', rec.get('pack') or '',
                rec['pack_kg'] if rec.get('pack_kg') is not None else Decimal(1),
                rec['stock'] if rec.get('stock') is not None else 0,
                rec['active'] if rec.get('active') is not None else True,
                rec.get('group_id'), rec.get('category') or 'Прочее', sort,
            ))
        execute_values(cur, f"INSERT INTO {schema}.products (id, name, full_name, article, code_1c, barcode, unit, "
                            f"manufacturer, dimensions, pack, pack_kg, stock, active, group_id, category, sort) VALUES %s",
                       data, page_size=500)

    if upd_rows:
        data = [(pid, *(rec.get(c) for c in UPD_COLS)) for pid, rec in upd_rows.items()]
        sets = ', '.join(f"{c} = COALESCE(v.{c}, p.{c})" for c in UPD_COLS)
        tpl = '(%s, ' + ', '.join(f"%s::{UPD_TYPES[c]}" for c in UPD_COLS) + ')'
        execute_values(cur, f"UPDATE {schema}.products p SET {sets}, updated_at = NOW() "
                            f"FROM (VALUES %s) AS v (id, {', '.join(UPD_COLS)}) WHERE p.id = v.id",
                       data, template=tpl, page_size=500)

    if price_map:
        execute_values(cur, f"INSERT INTO {schema}.product_prices (product_id, price_type_id, price) VALUES %s "
                            f"ON CONFLICT (product_id, price_type_id) DO UPDATE SET price = EXCLUDED.price, updated_at = NOW()",
                       [(pid, tid, val) for (pid, tid), val in price_map.items()], page_size=500)
        ids = list({pid for pid, _ in price_map})
        _q(cur, "UPDATE {S}.products p SET price = pp.price, updated_at = NOW() FROM {S}.product_prices pp "
                "JOIN {S}.price_types t ON t.id = pp.price_type_id AND t.is_main "
                "WHERE pp.product_id = p.id AND p.id = ANY(%s)", (ids,))

    created, updated = len(new_rows), len(upd_rows)
    return 200, {'ok': True, 'created': created, 'updated': updated, 'errors': errors[:50]}
