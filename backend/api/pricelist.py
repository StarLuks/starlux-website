"""Прайс-лист клиента в Excel с его ценами и миниатюрами товаров."""
import base64
import io
import os
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from decimal import Decimal

SCHEMA = os.environ.get('MAIN_DB_SCHEMA', 't_p16770056_starlux_website')
THUMB = 64


def _thumb(url):
    from PIL import Image
    try:
        with urllib.request.urlopen(url, timeout=3) as r:
            img = Image.open(io.BytesIO(r.read()))
            img.thumbnail((THUMB, THUMB))
            if img.mode not in ('RGB', 'RGBA'):
                img = img.convert('RGBA')
            buf = io.BytesIO()
            img.save(buf, 'PNG')
            buf.seek(0)
            return buf
    except Exception:
        return None


def build_price_list(cur, price_type_id, company):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    cur.execute(
        f"SELECT p.id, p.name, p.article, p.unit, p.manufacturer, p.pack, p.pack_kg, p.stock, p.category, "
        f"COALESCE(pp.price, p.price) AS price, "
        f"(SELECT i.url FROM {SCHEMA}.product_images i WHERE i.product_id = p.id "
        f" ORDER BY i.is_main DESC, i.sort, i.id LIMIT 1) AS image "
        f"FROM {SCHEMA}.products p LEFT JOIN {SCHEMA}.product_prices pp "
        f"ON pp.product_id = p.id AND pp.price_type_id = %s "
        f"WHERE p.active ORDER BY p.category, p.sort, p.name", (price_type_id or 0,))
    rows = cur.fetchall()

    urls = list({r['image'] for r in rows if r['image']})
    with ThreadPoolExecutor(max_workers=8) as ex:
        thumbs = dict(zip(urls, ex.map(_thumb, urls)))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Прайс-лист'
    blue = '1D4F9C'
    thin = Side(style='thin', color='D9DEE7')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws['A1'] = 'Прайс-лист СтарЛюкс'
    ws['A1'].font = Font(bold=True, size=14, color=blue)
    ws['A2'] = f"{company or ''} · на {datetime.now().strftime('%d.%m.%Y')}"
    ws['A2'].font = Font(color='6B7280')

    header = ['Фото', 'Наименование', 'Артикул', 'Производитель', 'Ед. изм.', 'Фасовка',
              'Цена за кг, ₽', 'Цена за ед., ₽', 'Остаток']
    widths = [11, 44, 14, 22, 10, 16, 14, 15, 10]
    hr = 4
    for i, (h, w) in enumerate(zip(header, widths), start=1):
        c = ws.cell(hr, i, h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=blue)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[hr].height = 30

    r = hr
    cat = None
    for p in rows:
        if p['category'] != cat:
            cat = p['category']
            r += 1
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(header))
            c = ws.cell(r, 1, cat or 'Прочее')
            c.font = Font(bold=True, color=blue)
            c.fill = PatternFill('solid', fgColor='EEF2F8')
            ws.row_dimensions[r].height = 22
        r += 1
        price = Decimal(p['price'] or 0)
        unit_price = round(price * Decimal(p['pack_kg'] or 1))
        vals = [None, p['name'], p['article'] or '', p['manufacturer'] or '', p['unit'] or 'кор.',
                p['pack'] or '', float(price), float(unit_price), p['stock']]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(r, i, v)
            c.border = border
            c.alignment = Alignment(vertical='center', wrap_text=i == 2,
                                    horizontal='right' if i >= 7 else ('center' if i == 5 else 'left'))
            if i in (7, 8):
                c.number_format = '#,##0.00' if i == 7 else '#,##0'
        ws.cell(r, 8).font = Font(bold=True)
        t = thumbs.get(p['image']) if p['image'] else None
        if t:
            img = XlImage(t)
            ws.add_image(img, f"A{r}")
            ws.row_dimensions[r].height = 52
        else:
            ws.row_dimensions[r].height = 22

    ws.freeze_panes = ws.cell(hr + 1, 1)
    ws.auto_filter.ref = f"A{hr}:{ws.cell(hr, len(header)).column_letter}{r}"
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()
